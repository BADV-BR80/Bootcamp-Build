"""The Python to Excel handoff."""

import openpyxl
import pytest

from bootcamp.emit import write_inputs_workbook
from bootcamp.validate import validate_table


@pytest.fixture
def workbook(tmp_path, hours_df, roster_df):
    tables = {
        "hours_entry": validate_table("hours_entry", hours_df),
        "roster": validate_table("roster", roster_df),
    }
    path = write_inputs_workbook(
        tables, tmp_path / "inputs.xlsx", client="_demo", sources={"hours_entry": [tmp_path / "hours.csv"]}
    )
    return openpyxl.load_workbook(path)


def test_one_sheet_per_table_plus_provenance(workbook):
    assert workbook.sheetnames == ["Provenance", "hours_entry", "roster"]


def test_headers_are_the_canonical_field_names(workbook):
    header = [c.value for c in workbook["hours_entry"][1]]
    assert header[:3] == ["entry_id", "date", "employee_ref"]


def test_rows_survive(workbook):
    assert workbook["hours_entry"].max_row == 4  # header + 3
    assert workbook["roster"].max_row == 3


def test_provenance_records_source_and_coverage(workbook):
    text = "\n".join(
        str(c.value) for row in workbook["Provenance"].iter_rows() for c in row if c.value
    )
    assert "_demo" in text
    assert "generic_csv" in text
    assert "2025-05-01 to 2025-05-02" in text
    assert "hours.csv" in text


def test_destination_directory_is_created(tmp_path, hours_df):
    tables = {"hours_entry": validate_table("hours_entry", hours_df)}
    out = write_inputs_workbook(tables, tmp_path / "deep" / "nested" / "inputs.xlsx", client="x")
    assert out.exists()
