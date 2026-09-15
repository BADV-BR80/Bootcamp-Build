"""Layer 2 → layer 3 handoff.

Writes the validated canonical tables to one workbook per client. That workbook
is the contract between the Python half of the build and the Excel half: one
sheet per canonical table, plus a provenance sheet recording what came from
where.

The handoff is a file a human can open on purpose. Everything the extract layer
did is visible in it, which is what makes the numbers defensible later.
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .canonical import TABLES

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")


def _autosize(worksheet, df: pd.DataFrame, *, max_width: int = 42) -> None:
    for idx, col in enumerate(df.columns, start=1):
        body = df[col].astype("string").fillna("")
        widest = max([len(str(col))] + [len(v) for v in body.head(500)]) + 2
        worksheet.column_dimensions[get_column_letter(idx)].width = min(widest, max_width)


def _style_header(worksheet, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = worksheet.cell(row=1, column=c)
        cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center")
    worksheet.freeze_panes = "A2"


def _provenance(tables: dict[str, pd.DataFrame], sources: dict[str, list[Path]]) -> pd.DataFrame:
    rows = []
    for name, df in tables.items():
        systems = (
            ", ".join(sorted(df["source_system"].dropna().astype(str).unique()))
            if "source_system" in df.columns and len(df)
            else ""
        )
        coverage = ""
        for date_col in ("date", "month", "period_start"):
            if date_col in df.columns and df[date_col].notna().any():
                lo = pd.to_datetime(df[date_col]).min().date()
                hi = pd.to_datetime(df[date_col]).max().date()
                coverage = f"{lo} to {hi}"
                break
        rows.append(
            {
                "canonical_table": name,
                "grain": TABLES[name].grain,
                "spec_ref": TABLES[name].spec_ref,
                "rows": len(df),
                "source_system": systems,
                "date_coverage": coverage,
                "source_files": "; ".join(p.name for p in sources.get(name, [])),
            }
        )
    return pd.DataFrame(rows)


def write_inputs_workbook(
    tables: dict[str, pd.DataFrame],
    destination: Path,
    *,
    client: str,
    sources: dict[str, list[Path]] | None = None,
) -> Path:
    """Write validated canonical tables to `destination` as one workbook."""
    destination.parent.mkdir(parents=True, exist_ok=True)
    sources = sources or {}
    provenance = _provenance(tables, sources)
    stamp = dt.datetime.now(dt.timezone.utc).strftime("%Y-%m-%d %H:%M UTC")

    with pd.ExcelWriter(destination, engine="openpyxl") as writer:
        header = pd.DataFrame(
            {
                "field": ["client", "generated", "canonical tables", "layer"],
                "value": [
                    client,
                    stamp,
                    ", ".join(tables),
                    "2 — canonical (see data-architecture.md)",
                ],
            }
        )
        header.to_excel(writer, sheet_name="Provenance", index=False, startrow=0)
        provenance.to_excel(writer, sheet_name="Provenance", index=False, startrow=len(header) + 2)
        ws = writer.sheets["Provenance"]
        _style_header(ws, len(header.columns))
        for c in range(1, len(provenance.columns) + 1):
            cell = ws.cell(row=len(header) + 3, column=c)
            cell.font = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
            cell.fill = HDR_FILL
        for col, width in zip("ABCDEFG", [22, 30, 26, 10, 16, 24, 30]):
            ws.column_dimensions[col].width = width

        for name, df in tables.items():
            sheet = name[:31]
            df.to_excel(writer, sheet_name=sheet, index=False)
            worksheet = writer.sheets[sheet]
            _style_header(worksheet, len(df.columns))
            _autosize(worksheet, df)

    return destination
